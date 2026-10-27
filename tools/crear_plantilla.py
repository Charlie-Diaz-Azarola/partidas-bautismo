"""Genera la plantilla en blanco de la certificación de partida de bautismo (Word).

Reproduce el formato del impreso «Certificación de partida de bautismo»:
columna izquierda (datos del archivo y notas marginales) y cuerpo a la derecha.
Los campos entre llaves, por ejemplo {nombre}, los sustituye la web al rellenar
la partida. Se puede abrir plantilla/plantilla_partida.docx en Word y retocar el
diseño, siempre que se conserven los campos.

Uso:  python tools/crear_plantilla.py
"""
import io
import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

DESTINO = Path(__file__).resolve().parent.parent / "plantilla" / "plantilla_partida.docx"

ANCHO_IZQ = Cm(5.6)
ANCHO_DER = Cm(12.4)

doc = Document()
seccion = doc.sections[0]
seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
seccion.top_margin = seccion.bottom_margin = Cm(1.5)
seccion.left_margin = seccion.right_margin = Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)

CAMPO = re.compile(r"(\{[a-z_0-9]+\})")


def parrafo(contenedor, partes, *, tam=11, fuente=None, alinear=None,
            antes=0, despues=4, sangria=0):
    """Añade un párrafo. `partes` es un texto o una lista de (negrita, texto).
    Lo que va entre llaves es un campo (se muestra punteado); el resto, texto fijo."""
    if isinstance(partes, str):
        partes = [(False, partes)]
    p = contenedor.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if sangria:
        p.paragraph_format.left_indent = Cm(sangria)
    if alinear:
        p.alignment = alinear
    for negrita, texto in partes:
        for trozo in CAMPO.split(texto):
            if not trozo:
                continue
            r = p.add_run(trozo)
            r.font.size = Pt(tam)
            if fuente:
                r.font.name = fuente
            if CAMPO.fullmatch(trozo):
                r.font.underline = WD_UNDERLINE.DOTTED
            else:
                r.bold = negrita
    return p


def borde_derecho(celda):
    bordes = OxmlElement("w:tcBorders")
    b = OxmlElement("w:right")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "0"), ("color", "808080")):
        b.set(qn(f"w:{k}"), v)
    bordes.append(b)
    celda._tc.get_or_add_tcPr().append(bordes)


# --- Título ----------------------------------------------------------------
titulo = parrafo(doc, [(True, "CERTIFICACIÓN DE PARTIDA DE BAUTISMO")], tam=17,
                 alinear=WD_ALIGN_PARAGRAPH.CENTER, antes=6, despues=10)
for r in titulo.runs:
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), "40")  # espaciado entre letras
    r._r.get_or_add_rPr().append(sp)
# Línea horizontal bajo el título que lo separa del resto del contenido
bordes_titulo = OxmlElement("w:pBdr")
linea = OxmlElement("w:bottom")
for k, v in (("val", "single"), ("sz", "8"), ("space", "6"), ("color", "000000")):
    linea.set(qn(f"w:{k}"), v)
bordes_titulo.append(linea)
titulo._p.get_or_add_pPr().insert(0, bordes_titulo)

# --- Cuerpo en dos columnas -------------------------------------------------
tabla = doc.add_table(rows=1, cols=2)
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
tabla.autofit = False
izq, der = tabla.rows[0].cells
izq.width, der.width = ANCHO_IZQ, ANCHO_DER
borde_derecho(izq)
for c in (izq, der):  # quita el párrafo vacío inicial de cada celda
    c._tc.remove(c.paragraphs[0]._p)

# Columna izquierda
for etiqueta, campo in (("Parroquia", "parroquia"), ("Población", "poblacion"),
                        ("Diócesis", "diocesis"), ("Provincia", "provincia")):
    parrafo(izq, f"{etiqueta} {{{campo}}}", tam=9, fuente="Arial", despues=9)
for etiqueta, campo in (("Libro", "libro"), ("Folio", "folio"), ("Núm.", "numero")):
    parrafo(izq, f"{etiqueta} {{{campo}}}", tam=9, fuente="Arial", despues=3)
parrafo(izq, [(True, "Notas marginales")], tam=10, fuente="Arial",
        alinear=WD_ALIGN_PARAGRAPH.CENTER, antes=18, despues=8)
parrafo(izq, "{notas_marginales}", tam=9, fuente="Arial")

# Columna derecha
SANG = 1.2  # sangría de las líneas PADRES / ABUELOS
parrafo(der, [(True, "Don "), (False, "{encargado},")], despues=6)
parrafo(der, "Encargado del Archivo Parroquial de {archivo_de}", despues=6)
parrafo(der, "Diócesis de {diocesis},", despues=6)
parrafo(der, [(True, "CERTIFICA: "),
              (False, "Que según consta del acta reseñada al margen, correspondiente "
                      "al libro de Bautismos,")])
parrafo(der, "D. {nombre},")
parrafo(der, "fue BAUTIZADO el día {bautismo_dia} de {bautismo_mes} de {bautismo_anio}")
parrafo(der, "Nació el día {nacimiento_dia} de {nacimiento_mes} de {nacimiento_anio},")
parrafo(der, "en la calle {calle}, n.º {calle_numero},")
parrafo(der, "siendo natural de {natural_de}, Diócesis de {natural_diocesis},")
parrafo(der, "Provincia de {natural_provincia}", despues=8)

for titulo_bloque, hombre, mujer in (("PADRES", "padre", "madre"),
                                     ("ABUELOS PATERNOS", "abuelo_paterno", "abuela_paterna"),
                                     ("ABUELOS MATERNOS", "abuelo_materno", "abuela_materna")):
    parrafo(der, f"{titulo_bloque}: D. {{{hombre}}},", sangria=SANG, despues=3)
    parrafo(der, f"natural de {{{hombre}_natural}},", despues=3)
    parrafo(der, f"y D.ª {{{mujer}}},", despues=3)
    parrafo(der, f"natural de {{{mujer}_natural}}", despues=8)

parrafo(der, "PADRINOS: {padrinos}", despues=8)
parrafo(der, "MINISTRO: {ministro}", despues=14)
parrafo(der, "{lugar_expedicion}, a {expedicion_dia} de {expedicion_mes} de {expedicion_anio}",
        alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
# Imagen de relleno transparente (misma proporción que la firma y el sello reales).
# La web la sustituye por la imagen real, que sirve el Worker (word/media/image1.png).
relleno = io.BytesIO()
Image.new("RGBA", (133, 100), (0, 0, 0, 0)).save(relleno, "PNG")
imagen = parrafo(der, "", alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
imagen.add_run().add_picture(relleno, width=Cm(3.6))
parrafo(der, "(Firma del Encargado del Archivo)", tam=8, fuente="Arial",
        alinear=WD_ALIGN_PARAGRAPH.RIGHT, despues=0)
parrafo(der, "(Sello)", tam=8, fuente="Arial", alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=14)

# Pie: para otras diócesis
parrafo(der, "(Para otras Diócesis)", tam=8, fuente="Arial", sangria=1.2, despues=2)
parrafo(der, "Obispado de {obispado}", tam=8, fuente="Arial", despues=2)
parrafo(der, "V.º B.º:", tam=8, fuente="Arial", sangria=2.5, despues=0)
parrafo(der, "El Vicario General,", tam=8, fuente="Arial", sangria=2.2, despues=0)

DESTINO.parent.mkdir(parents=True, exist_ok=True)
doc.save(DESTINO)
print("Plantilla creada en", DESTINO)
